
"""
Ejercicio 05 - Book Manager.

Genera un reporte Excel a partir del CSV del Ejercicio 04.

Columnas:
- Libro
- Precio interno
- Precio web
- Diferencia
- Fecha de extracción
"""

import csv
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


# ============================================================
# CONVERTIR VALORES NUMÉRICOS
# ============================================================

def convertir_numero(valor):
    """
    Convierte un valor del CSV a número.

    Si está vacío o no es válido, devuelve None.
    """

    if valor is None or str(valor).strip() == "":
        return None

    try:
        numero = Decimal(str(valor).strip())

        if not numero.is_finite():
            return None

        return float(numero)

    except (InvalidOperation, ValueError, TypeError):
        return None


# ============================================================
# CONVERTIR FECHA
# ============================================================

def convertir_fecha(valor):
    """
    Convierte la fecha de ejecución del CSV
    a un valor de fecha compatible con Excel.
    """

    if not valor:
        return None

    try:
        return datetime.fromisoformat(valor)

    except (ValueError, TypeError):
        return None


# ============================================================
# LEER CSV DEL EJERCICIO 04
# ============================================================

def leer_csv(ruta_csv):
    """
    Lee el CSV generado por el Ejercicio 04.
    """

    ruta_csv = Path(ruta_csv)

    if not ruta_csv.is_file():
        raise FileNotFoundError(
            f"No se encontró el CSV: {ruta_csv}"
        )

    with ruta_csv.open(
        "r",
        encoding="utf-8-sig",
        newline=""
    ) as archivo:

        lector = csv.DictReader(
            archivo,
            delimiter=";"
        )

        columnas_requeridas = {
            "titulo",
            "precio_interno_ars",
            "precio_web_ars",
            "diferencia_ars",
            "fecha_ejecucion",
        }

        if not columnas_requeridas.issubset(
            set(lector.fieldnames or [])
        ):
            raise ValueError(
                "El CSV no contiene las columnas necesarias."
            )

        return list(lector)


# ============================================================
# GENERAR REPORTE EXCEL
# ============================================================

def generar_excel(ruta_csv, ruta_excel=None):
    """
    Genera un archivo Excel con las cinco columnas
    solicitadas en el Ejercicio 05.

    Cuando no existe precio web:
    - Precio web = Sin datos encontrados
    - Diferencia = Sin datos encontrados
    """

    ruta_csv = Path(ruta_csv)

    registros = leer_csv(ruta_csv)

    if ruta_excel is None:
        ruta_excel = ruta_csv.with_name(
            ruta_csv.stem.replace(
                "alertas_precios",
                "reporte_precios"
            ) + ".xlsx"
        )

    ruta_excel = Path(ruta_excel)

    ruta_excel.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    # Crear libro de Excel
    libro_excel = Workbook()

    hoja = libro_excel.active
    hoja.title = "Comparación de precios"

    # Encabezados solicitados
    encabezados = [
        "Libro",
        "Precio interno",
        "Precio web",
        "Diferencia",
        "Fecha de extracción",
    ]

    hoja.append(encabezados)

    # ========================================================
    # CARGAR REGISTROS
    # ========================================================

    for registro in registros:

        titulo = registro["titulo"]

        precio_interno = convertir_numero(
            registro["precio_interno_ars"]
        )

        precio_web = convertir_numero(
            registro["precio_web_ars"]
        )

        diferencia = convertir_numero(
            registro["diferencia_ars"]
        )

        # Cuando no se encontró un precio válido en la web,
        # mostrar un mensaje en lugar de cero.
        if precio_web is None:

            precio_web = "Sin datos encontrados"
            diferencia = "Sin datos encontrados"

        fecha = convertir_fecha(
            registro["fecha_ejecucion"]
        )

        hoja.append([
            titulo,
            precio_interno,
            precio_web,
            diferencia,
            fecha,
        ])

    # ========================================================
    # FORMATO DE ENCABEZADOS
    # ========================================================

    color_encabezado = PatternFill(
        fill_type="solid",
        fgColor="17365D"
    )

    fuente_encabezado = Font(
        color="FFFFFF",
        bold=True
    )

    for celda in hoja[1]:

        celda.fill = color_encabezado
        celda.font = fuente_encabezado

        celda.alignment = Alignment(
            horizontal="center"
        )

    # ========================================================
    # FORMATO DE VALORES
    # ========================================================

    for fila in hoja.iter_rows(
        min_row=2,
        max_row=hoja.max_row
    ):

        # Precio interno, web y diferencia
        for celda in fila[1:4]:

            # Aplicar formato moneda únicamente a números.
            if isinstance(celda.value, (int, float)):

                celda.number_format = (
                    '"$" #,##0.00;[Red]("$" #,##0.00)'
                )

        # Fecha de extracción
        fila[4].number_format = "DD/MM/YYYY HH:MM"

    # ========================================================
    # ANCHO DE COLUMNAS
    # ========================================================

    anchos = {
        "A": 40,
        "B": 20,
        "C": 25,
        "D": 25,
        "E": 24,
    }

    for columna, ancho in anchos.items():
        hoja.column_dimensions[columna].width = ancho

    # Filtros
    hoja.auto_filter.ref = hoja.dimensions

    # Mantener encabezados visibles
    hoja.freeze_panes = "A2"

    # ========================================================
    # GUARDAR EXCEL
    # ========================================================

    libro_excel.save(ruta_excel)

    print("\n=== REPORTE EXCEL GENERADO ===")
    print("Registros:", len(registros))
    print("Archivo:", ruta_excel)

    return ruta_excel


# ============================================================
# EJECUCIÓN DIRECTA
# ============================================================

if __name__ == "__main__":
    print(
        "Este módulo debe ejecutarse indicando "
        "el CSV del Ejercicio 04."
    )

